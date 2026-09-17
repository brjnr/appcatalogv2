"""Tests for new features: server lifecycle, import/export, integrations, security."""
import io
import os
import time

import pytest
import requests
from openpyxl import Workbook, load_workbook

BASE_URL = (
    os.environ.get("APPCATALOG_PREVIEW_URL")
    or os.environ.get("BACKEND_URL")
    or "https://sast-review.preview.emergentagent.com"
)

ADMIN = {"email": "admin", "password": "Admin!Passw0rd_1"}
USER = {"email": "john.doe", "password": "User!Passw0rd_1"}


def _login(payload):
    s = requests.Session()
    r = s.post(f"{BASE_URL}/api/auth/login", json=payload, timeout=15)
    assert r.status_code == 200, r.text
    return s


@pytest.fixture(scope="module")
def admin_client():
    return _login(ADMIN)


# --------------------- Security regressions

def test_status_endpoint_returns_404():
    r = requests.get(f"{BASE_URL}/api/status", timeout=10)
    assert r.status_code == 404


def test_login_by_email_and_username():
    s1 = requests.Session()
    r1 = s1.post(f"{BASE_URL}/api/auth/login", json={"email": "admin@corp.com", "password": "Admin!Passw0rd_1"}, timeout=15)
    assert r1.status_code == 200
    s2 = requests.Session()
    r2 = s2.post(f"{BASE_URL}/api/auth/login", json={"email": "admin", "password": "Admin!Passw0rd_1"}, timeout=15)
    assert r2.status_code == 200


def test_brute_force_lockout_returns_429():
    # Use a throw-away identifier so we don't lock real accounts unnecessarily
    ident = f"noone_{int(time.time())}@corp.com"
    codes = []
    for _ in range(10):
        r = requests.post(f"{BASE_URL}/api/auth/login", json={"email": ident, "password": "wrong"}, timeout=10)
        codes.append(r.status_code)
        if r.status_code == 429:
            break
    # After ~5-8 attempts, a 429 should appear
    assert 429 in codes, f"expected 429 in {codes}"


# --------------------- Catalog regression

def test_admin_sees_min_36_apps(admin_client):
    r = admin_client.get(f"{BASE_URL}/api/apps", timeout=15)
    assert r.status_code == 200
    body = r.json()
    apps = body.get("items", body) if isinstance(body, dict) else body
    assert len(apps) >= 36, f"admin should see >=36 apps, got {len(apps)}"


def test_normal_user_sees_exactly_11_apps():
    s = _login(USER)
    r = s.get(f"{BASE_URL}/api/apps", timeout=15)
    assert r.status_code == 200
    body = r.json()
    apps = body.get("items", body) if isinstance(body, dict) else body
    assert len(apps) == 11, f"john.doe should see exactly 11 apps, got {len(apps)}"


# --------------------- Server lifecycle fields

def test_server_lifecycle_fields_persist(admin_client):
    payload = {
        "name": f"TEST_lifecycle_{int(time.time())}",
        "status": "Active",
        "date_created": "2024-06-15",
    }
    r = admin_client.post(f"{BASE_URL}/api/servers", json=payload, timeout=15)
    assert r.status_code == 201, r.text
    server = r.json()
    sid = server["id"]
    assert server["date_created"] == "2024-06-15"

    # GET verifies persistence
    r = admin_client.get(f"{BASE_URL}/api/servers/{sid}", timeout=15)
    assert r.json()["date_created"] == "2024-06-15"

    # Update to Decommissioned with dates + notes
    upd = {
        "status": "Decommissioned",
        "date_decommissioned": "2025-01-20",
        "decommission_notes": "Retired per DR migration",
    }
    r = admin_client.put(f"{BASE_URL}/api/servers/{sid}", json=upd, timeout=15)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "Decommissioned"
    assert body["date_decommissioned"] == "2025-01-20"
    assert body["decommission_notes"] == "Retired per DR migration"

    # Re-fetch
    r = admin_client.get(f"{BASE_URL}/api/servers/{sid}", timeout=15)
    b = r.json()
    assert b["date_created"] == "2024-06-15"
    assert b["date_decommissioned"] == "2025-01-20"
    assert b["decommission_notes"] == "Retired per DR migration"

    # cleanup
    admin_client.delete(f"{BASE_URL}/api/servers/{sid}", timeout=15)


# --------------------- Import / Export

def test_servers_template_download(admin_client):
    r = admin_client.get(f"{BASE_URL}/api/exports/servers/template", timeout=20)
    assert r.status_code == 200, r.text
    assert "spreadsheet" in r.headers.get("content-type", "")
    wb = load_workbook(io.BytesIO(r.content), read_only=True)
    ws = wb.active
    headers = [c.value for c in next(ws.iter_rows(max_row=1))]
    assert "date_created" in headers
    assert "date_decommissioned" in headers
    assert "decommission_notes" in headers


def test_import_csv_missing_name_returns_error_row(admin_client):
    csv_body = "name,hostname,ip_address\n,orphan.example,10.0.0.99\n"
    files = {"file": ("servers.csv", csv_body.encode(), "text/csv")}
    r = admin_client.post(f"{BASE_URL}/api/imports/servers", files=files, timeout=20)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["summary"]["errors"] >= 1, body
    assert body["results"][0]["status"] == "error"


def test_all_four_entity_templates(admin_client):
    for entity in ["categories", "applications", "servers", "pics"]:
        r = admin_client.get(f"{BASE_URL}/api/exports/{entity}/template", timeout=20)
        assert r.status_code == 200, f"{entity}: {r.text[:200]}"


# --------------------- Integrations

def test_list_integrations_has_5_connectors(admin_client):
    r = admin_client.get(f"{BASE_URL}/api/integrations", timeout=15)
    assert r.status_code == 200, r.text
    kinds = {item["kind"] for item in r.json()}
    assert kinds == {"smtp", "jira", "teams", "telegram", "vcenter"}, kinds


def test_smtp_password_masked_and_blank_preserves_value(admin_client):
    payload = {
        "enabled": True,
        "config": {
            "host": "smtp.example.com",
            "port": 587,
            "username": "svc@example.com",
            "password": "SuperSecret1234",
            "from_address": "no-reply@example.com",
            "use_tls": True,
        },
    }
    r = admin_client.put(f"{BASE_URL}/api/integrations/smtp", json=payload, timeout=15)
    assert r.status_code == 200, r.text
    masked = r.json()["config"]["password"]
    assert masked != "SuperSecret1234"
    assert "1234" in masked or "•" in masked or "*" in masked, masked

    # Refetch — still masked
    r = admin_client.get(f"{BASE_URL}/api/integrations/smtp", timeout=15)
    assert r.status_code == 200
    m2 = r.json()["config"]["password"]
    assert m2 != "SuperSecret1234"
    assert r.json()["has_secret"] is True

    # Save with blank password — should keep existing
    payload2 = dict(payload)
    payload2["config"] = dict(payload["config"])
    payload2["config"]["password"] = ""
    r = admin_client.put(f"{BASE_URL}/api/integrations/smtp", json=payload2, timeout=15)
    assert r.status_code == 200
    assert r.json()["has_secret"] is True, "blank password must not clear stored secret"


# --------------------- Security controls

def test_svg_icon_upload_returns_415(admin_client):
    # Try uploading an SVG as a category icon → should be rejected 415
    # First get a category id
    r = admin_client.get(f"{BASE_URL}/api/categories", timeout=15)
    assert r.status_code == 200
    cat = r.json()[0]
    svg = b'<svg xmlns="http://www.w3.org/2000/svg"><rect/></svg>'
    files = {"file": ("icon.svg", svg, "image/svg+xml")}
    r = admin_client.post(f"{BASE_URL}/api/categories/{cat['id']}/icon", files=files, timeout=15)
    assert r.status_code in (415, 400), f"expected 415/400 got {r.status_code}: {r.text[:200]}"


def test_favorite_spam_is_idempotent(admin_client):
    r = admin_client.get(f"{BASE_URL}/api/apps", timeout=15)
    body = r.json()
    apps = body.get("items", body) if isinstance(body, dict) else body
    app_id = apps[0]["id"]
    # Unfavorite first to have a known state
    admin_client.post(f"{BASE_URL}/api/apps/{app_id}/favorite", json={"favorite": False}, timeout=15)
    baseline = admin_client.get(f"{BASE_URL}/api/apps/{app_id}", timeout=15).json().get("favorite_count", 0)
    # Spam favorite 5 times → must upsert idempotently
    for _ in range(5):
        rf = admin_client.post(f"{BASE_URL}/api/apps/{app_id}/favorite", json={"favorite": True}, timeout=15)
        assert rf.status_code == 200
    after = admin_client.get(f"{BASE_URL}/api/apps/{app_id}", timeout=15).json().get("favorite_count", 0)
    assert after - baseline == 1, f"expected +1 from single user, baseline={baseline} after={after}"
    # cleanup
    admin_client.post(f"{BASE_URL}/api/apps/{app_id}/favorite", json={"favorite": False}, timeout=15)


# --------------------- Password policy on user create

def test_weak_password_rejected_on_add_user(admin_client):
    r = admin_client.post(
        f"{BASE_URL}/api/users",
        json={
            "username": "TEST_weak",
            "email": "TEST_weak@corp.com",
            "full_name": "Weak Test",
            "role": "user",
            "password": "short",
            "password_confirm": "short",
        },
        timeout=15,
    )
    assert r.status_code in (400, 422), r.text


def test_password_confirm_mismatch_rejected(admin_client):
    r = admin_client.post(
        f"{BASE_URL}/api/users",
        json={
            "username": "TEST_mismatch",
            "email": "TEST_mismatch@corp.com",
            "full_name": "Mismatch Test",
            "role": "user",
            "password": "Good!Passw0rd_1",
            "password_confirm": "Different!Passw0rd_1",
        },
        timeout=15,
    )
    assert r.status_code in (400, 422), r.text
